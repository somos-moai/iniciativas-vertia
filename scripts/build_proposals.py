"""Build separate formal proposals with the authoritative Codex document runtime.

Run with bundled Python. Pandoc follows the supplied MoAI Word reference.
"""
from pathlib import Path
import base64, html, json, re, shutil, subprocess, sys, textwrap
from decimal import Decimal
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'deliverables'
OUT.mkdir(exist_ok=True)
FIN = json.loads((ROOT/'docs/commercial/financials.json').read_text(encoding='utf-8-sig'))
REF = Path(r'C:\Users\cueva\OneDrive\Escritorio\Presupuestador MoAI\assets\moai-reference.docx')
if not REF.exists():
    candidates = list(REF.parent.parent.rglob('moai-reference.docx'))
    REF = candidates[0] if candidates else REF
PANDOC=shutil.which('pandoc') or r'C:\Users\cueva\AppData\Local\Pandoc\pandoc.exe'

def money(value,places=4):
    return f'{value:,.{places}f}'.replace(',','~').replace('.',',').replace('~','.')

def source_md(project):
    path=ROOT/'docs/commercial'/f"{project['id']}.md"
    source=path.read_text(encoding='utf-8-sig')
    # The canonical ledger replaces the draft tables, avoiding duplicated pricing.
    source=re.sub(r'\n## (?:Inversión|Calendario de pagos)[^\n]*\n.*?(?=\n## |\Z)','\n',source,flags=re.S)
    # Canonical appendix, generated from the reconciled payment ledger.
    source += '\n\n## Cuadro financiero contractual\n\n'
    source += '| Concepto | USD de referencia | UF de contratación |\n|---|---:|---:|\n'
    for label,field in [('Honorarios netos','net'),('IVA del 19 %','vat'),('Total con IVA','gross')]:
        source += f"| {label} | {money(project['usd_'+field],2)} | {money(project['uf_'+field])} |\n"
    source += '\n| Pago | Día propuesto | Neto UF | IVA UF referencial | Total UF |\n|---|---:|---:|---:|---:|\n'
    for pay in project['payments']:
        source += f"| {pay['description']} | {pay['due_day']} | {money(pay['uf_net'])} | {money(pay['uf_vat'])} | {money(pay['uf_gross'])} |\n"
    source += '\nLas seis cuotas mensuales financian el desarrollo. El anticipo corresponde al 40 % y el saldo al 60 %. La última cuota concilia el redondeo de cuatro decimales. Los días 30 a 180 desde el inicio son una condición de calendario propuesta; las fechas se fijan en la orden de compra.\n'
    source += '\nLa UF queda fijada como unidad contractual al emitir esta propuesta. El USD se muestra únicamente como referencia de origen; no se aplica un segundo reajuste por dólar. La factura se emite en CLP según la UF del vencimiento acordado, calculando el IVA sobre la base en CLP. Los importes de IVA en UF son referenciales.\n'
    source += '\nReferencias consultadas el 15 de octubre de 2026: UF CLP 41.162,73 y dólar observado CLP 987,03. [Banco Central de Chile](https://si3.bcentral.cl/indicadoressiete/secure/IndicadoresDiarios.aspx?Idioma=es-CL), [UF SII](https://www.sii.cl/valores_y_fechas/uf/uf2026.htm) e [IVA SII](https://www.sii.cl/preguntas_frecuentes/impuestos_mensuales/001_130_0572.htm). Vigencia comercial propuesta de 30 días.\n'
    return source

def flow(project):
    steps=['Identificar','Priorizar','Aprobar','Contactar','Registrar'] if project['id']=='prospectador' else ['Solicitar','Revisar','Aprobar','Firmar','Controlar']
    image=Image.new('RGB',(1500,180),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf',29)
    color='#0A2240' if project['id']=='prospectador' else '#103B58'
    for i,label in enumerate(steps):
        left=10+i*300;draw.rounded_rectangle((left,25,left+270,140),radius=10,fill=color)
        draw.text((left+135,82),label,font=font,fill='white',anchor='mm')
        if i<4:draw.line((left+274,82,left+294,82),fill='#647787',width=4);draw.polygon([(left+294,82),(left+286,76),(left+286,88)],fill='#647787')
    target=OUT/(project['id']+'-flujo.png');image.save(target)
    return target

CSS='''
:root{--ink:#103b58;--accent:#b8ee32;--paper:#f2f5f7;--muted:#61717e}*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:var(--paper);color:#223647;font-family:Inter,"Segoe UI",Arial,sans-serif;line-height:1.65}.rail{position:fixed;inset:0 auto 0 0;width:260px;padding:35px 28px;background:var(--ink);color:white;overflow:auto}.rail img{max-width:175px;max-height:54px;object-fit:contain;margin:0 0 24px}.rail p{font-size:12px;letter-spacing:1.4px;text-transform:uppercase;color:#b9cbd6}.rail a{display:block;color:#d4e1e9;text-decoration:none;font-size:13px;margin:0 0 12px}.rail a:hover{color:var(--accent)}main{max-width:1180px;margin-left:260px;padding:58px 65px 80px}.mast{font-size:12px;letter-spacing:2px;color:var(--muted);text-transform:uppercase;margin-bottom:16px}h1{font-size:46px;line-height:1.14;letter-spacing:-1.4px;margin:0 0 22px;color:var(--ink);max-width:760px}h2{font-size:25px;color:var(--ink);margin:0 0 20px;line-height:1.25}h3{font-size:18px;line-height:1.3;color:var(--ink);margin:28px 0 12px}p{margin:0 0 17px}section{background:white;border:1px solid #dce5eb;border-radius:16px;padding:38px;margin:25px 0;scroll-margin-top:24px}section:first-of-type{border-top:4px solid var(--accent)}.summary{display:grid;grid-template-columns:repeat(3,1fr);gap:16px;margin:30px 0}.summary div{padding:22px;background:var(--ink);color:white;border-radius:12px}.summary span{display:block;font-size:11px;letter-spacing:1px;text-transform:uppercase;color:#bfd0da}.summary strong{display:block;font-size:25px;letter-spacing:-.5px}.summary small{font-size:12px;color:#cfdae1}.actions{display:flex;gap:12px;flex-wrap:wrap;margin:22px 0 38px}.actions a,.actions button{border:1px solid #cbd8df;background:white;color:var(--ink);padding:10px 16px;border-radius:8px;text-decoration:none;cursor:pointer;font:inherit;font-size:13px}.actions a:first-child{background:var(--ink);color:white}.section-kicker{font-size:11px;letter-spacing:2px;text-transform:uppercase;color:var(--muted);display:block;margin-bottom:8px}table{border-collapse:collapse;width:100%;font-size:13px;margin:24px 0;display:table}th{background:var(--ink);color:white;text-align:left}td,th{padding:12px;border:1px solid #d9e1e7;vertical-align:top}tr:nth-child(even) td{background:#f5f8fa}ul,ol{padding-left:22px}li{margin:8px 0}a{color:#146c9b}footer{font-size:12px;color:var(--muted);padding-top:20px}blockquote{margin:18px 0;padding:0 0 0 18px;border-left:2px solid #d5dfe4;color:#536b7b}@media(max-width:1000px){.rail{position:static;width:100%;height:auto;padding:20px 25px}.rail nav{display:none}.rail img{margin:0;max-height:35px}.rail p{display:none}main{margin:0;padding:35px 25px}.summary{grid-template-columns:1fr}h1{font-size:34px}section{padding:24px;overflow-x:auto}}@media print{body{background:white;font-size:10.5pt}.rail,.actions{display:none}main{margin:0;padding:0;max-width:none}section{border:0;border-radius:0;padding:0;margin:22px 0;break-inside:auto}.summary{grid-template-columns:repeat(3,1fr)}.summary div{border:1px solid #ddd;background:white;color:#000}.summary span,.summary small{color:#555}h1{font-size:26pt}h2{font-size:16pt;break-after:avoid}h3{break-after:avoid}tr{break-inside:avoid}a{color:inherit;text-decoration:none}@page{size:letter;margin:20mm}}
'''

def build_html(project,md):
    ident=project['id']; title='Prospectador comercial' if ident=='prospectador' else 'Gestión Contractual'
    tmp=OUT/f'{ident}.md'; tmp.write_text(md,encoding='utf-8')
    result=subprocess.run([PANDOC,str(tmp),'-f','markdown','-t','html5'],capture_output=True,text=True,encoding='utf-8',check=True).stdout
    result=re.sub(r'^<h1[^>]*>.*?</h1>\s*','',result,count=1,flags=re.S)
    flow64=base64.b64encode(flow(project).read_bytes()).decode()
    result=re.sub(r'(<h2[^>]*>Un día con la solución</h2>)',r'\1'+f'<img alt="Flujo de trabajo de la solución" style="width:100%;height:auto;margin:12px 0" src="data:image/png;base64,{flow64}">',result)
    parts=re.split(r'(<h2[^>]*>.*?</h2>)',result,flags=re.S)
    sections=[]; nav=[]
    if parts[0].strip(): sections.append('<section><span class="section-kicker">Propuesta comercial independiente</span>'+parts[0]+'</section>')
    for i in range(1,len(parts),2):
        heading=re.sub('<[^>]+>','',parts[i]); sid=f'seccion-{i}'
        nav.append(f'<a href="#{sid}">{html.escape(heading)}</a>')
        sections.append(f'<section id="{sid}"><span class="section-kicker">{(i+1)//2:02d} / {title}</span>{parts[i]}{parts[i+1] if i+1<len(parts) else ""}</section>')
    logo=ROOT/'frontend/public/brand'/('global-oficial.webp' if ident=='prospectador' else 'vertia.webp')
    brand64=base64.b64encode(logo.read_bytes()).decode()
    pdf=f'{ident}.pdf'; docx=f'{ident}.docx'
    theme=':root{--ink:#0a2240;--accent:#ce0e2d}' if ident=='prospectador' else ''
    page=f'''<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Propuesta {title} · MoAI para Vertia</title><style>{CSS}{theme}</style><aside class="rail"><img alt="{'Global Vía Pública' if ident=='prospectador' else 'Vertia Group'}" src="data:image/webp;base64,{brand64}"><p>MoAI · Propuesta 2026</p><nav>{''.join(nav)}</nav></aside><main><div class="mast">MoAI / Vertia Group / 15 de octubre de 2026</div><h1>{title}</h1><p>Implementación por etapas, alcance verificable y operación regional.</p><div class="summary"><div><span>Inversión neta</span><strong>{money(project['uf_net'])} UF</strong><small>USD {money(project['usd_net'],0)} de referencia</small></div><div><span>Forma de pago</span><strong>40 % + 6 cuotas</strong><small>IVA 19 % adicional</small></div><div><span>Modelo de entrega</span><strong>MVP + consolidación</strong><small>Servicios externos a cargo de Vertia</small></div></div><div class="actions"><a href="{pdf}" download>Descargar propuesta PDF</a><a href="{docx}" download>Versión editable Word</a><button onclick="window.print()">Imprimir</button></div>{''.join(sections)}<footer>MoAI · somos@mo-ai.cl · Documento de propuesta, sujeto a formalización de las condiciones comerciales.</footer></main></html>'''
    (OUT/f'{ident}.html').write_text(page,encoding='utf-8')

def style_docx(path,project):
    doc=Document(path)
    for section in doc.sections:
        section.page_width=Inches(8.5); section.page_height=Inches(11)
        section.top_margin=Inches(.7); section.bottom_margin=Inches(.65)
        section.left_margin=Inches(.8); section.right_margin=Inches(.8)
    for style in doc.styles:
        if style.type==1:
            style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal']; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',25),('Heading 1',19),('Heading 2',15),('Heading 3',12)]:
        st=doc.styles[name];st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
        if st.element.pPr is not None:
            for el in list(st.element.pPr):
                if el.tag==qn('w:pBdr'):st.element.pPr.remove(el)
    for p in doc.paragraphs:
        if p.text.strip():
            p.style=doc.styles['Title']; break
    for p in list(doc.paragraphs):
        if p.text.strip()=='Cuadro financiero contractual':
            p.paragraph_format.page_break_before=True
        if p.text.strip()=='Un día con la solución':
            img=doc.add_paragraph();img.add_run().add_picture(str(flow(project)),width=Inches(6.75))
            p._p.addnext(img._p)
    # Official co-branding is kept compact in the first-page header.
    brand=Image.new('RGB',(1400,150),'white')
    moai=Path(r'C:\Users\cueva\OneDrive\Escritorio\Presupuestador MoAI\.claude\skills\propuesta-moai\assets\moai-logo-negro.png')
    if moai.exists():
        mark=Image.open(moai).convert('RGBA');mark.thumbnail((300,115));brand.paste(mark,(0,15),mark)
    client=Image.open(ROOT/'frontend/public/brand'/('global-oficial.webp' if project['id']=='prospectador' else 'vertia.webp')).convert('RGBA')
    client.thumbnail((370,110))
    ImageDraw.Draw(brand).rectangle((1000,0,1399,149),fill='#0A2240' if project['id']=='prospectador' else '#103B58')
    brand.paste(client,(1020,15),client)
    brandpath=OUT/(project['id']+'-marcas.png');brand.save(brandpath)
    for table in doc.tables:
        table.autofit=True
        for i,row in enumerate(table.rows):
            if i==0:
                rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
            for cell in row.cells:
                cell.vertical_alignment=1
                tcpr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                tcpr.append(borders)
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'103B58' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');margins.append(el)
                tcpr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    for run in p.runs:
                        run.font.size=Pt(9.5);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    for section in doc.sections:
        hp=section.header.paragraphs[0]; hp.text='MoAI  |  Propuesta para Vertia Group'
        for run in hp.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
        fp=section.footer.paragraphs[0];fp.text='15 octubre 2026  ·  somos@mo-ai.cl                                      '
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
        section.different_first_page_header_footer=True
        first=section.first_page_header.paragraphs[0];first.clear();first.add_run().add_picture(str(brandpath),width=Inches(6.85))
        firstfooter=section.first_page_footer.paragraphs[0];firstfooter.text='15 octubre 2026  ·  somos@mo-ai.cl'
    doc.save(path)

for project in FIN['projects']:
    ident=project['id'];md=source_md(project)
    build_html(project,md)
    args=[PANDOC,str(OUT/f'{ident}.md'),'-o',str(OUT/f'{ident}.docx')]
    if REF.exists(): args+=['--reference-doc',str(REF)]
    subprocess.run(args,check=True)
    style_docx(OUT/f'{ident}.docx',project)
    print(f'Built {ident}: HTML, Markdown, DOCX',flush=True)
(OUT/'PAGOS_UF_v1.json').write_text(json.dumps(FIN,ensure_ascii=False,indent=2),encoding='utf-8')
